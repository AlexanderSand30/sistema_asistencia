from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generar_excel(titulo, encabezados, filas, nombre_hoja="Reporte"):
    """
    titulo: texto de la fila superior (combinada)
    encabezados: lista de strings para la fila de cabecera
    filas: lista de listas/tuplas con los valores de cada fila
    """
    wb = Workbook()
    ws = wb.active
    ws.title = nombre_hoja

    NAVY = "2C3E50"
    BLANCO = "FFFFFF"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(encabezados))
    celda_titulo = ws.cell(row=1, column=1, value=titulo)
    celda_titulo.font = Font(bold=True, size=14, color=BLANCO)
    celda_titulo.fill = PatternFill("solid", fgColor=NAVY)
    celda_titulo.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    for col_idx, encabezado in enumerate(encabezados, start=1):
        celda = ws.cell(row=2, column=col_idx, value=encabezado)
        celda.font = Font(bold=True, color=BLANCO)
        celda.fill = PatternFill("solid", fgColor=NAVY)
        celda.alignment = Alignment(horizontal="center")

    borde_fino = Border(bottom=Side(style="thin", color="DDDDDD"))
    for fila_idx, fila in enumerate(filas, start=3):
        for col_idx, valor in enumerate(fila, start=1):
            celda = ws.cell(row=fila_idx, column=col_idx, value=valor)
            celda.border = borde_fino
            if fila_idx % 2 == 0:
                celda.fill = PatternFill("solid", fgColor="F8F9FA")

    for col_idx, encabezado in enumerate(encabezados, start=1):
        max_len = len(str(encabezado))
        for fila in filas:
            valor = fila[col_idx - 1]
            max_len = max(max_len, len(str(valor)) if valor is not None else 0)
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()